import os
import json
import re
import base64
import traceback
import fitz  # PyMuPDF
from django.core.files.base import ContentFile
from django.conf import settings
from .models import Quiz, Question, Choice

try:
    from docx import Document
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    from google import genai
    from google.genai import types as genai_types
    HAS_GENAI = True
except ImportError:
    HAS_GENAI = False
    print("WARNING: google-genai package not installed. AI extraction disabled.")


class QuizExtractionService:
    def __init__(self):
        self.api_key = os.getenv("GEMINI_API_KEY", "AIzaSyCSztpNphwSaIOEtpsmgrS0kpm8zQVKOsU")
        if HAS_GENAI and self.api_key:
            self.client = genai.Client(api_key=self.api_key)
            self.model_name = "gemini-2.5-flash-lite"
        else:
            self.client = None
            self.model_name = None

    EXTRACTION_PROMPT = """
    You are an expert at extracting quiz questions from academic papers. 
    Your task is to convert the following text into a structured JSON array of multiple-choice questions.

    ### CORE RULES:
    1. **STRICT TEXT ADHERENCE**: Extract ONLY what is present in the text provided below. 
    2. **NO EXTERNAL KNOWLEDGE**: Do NOT use your internal training data to "guess" questions or equations for specific exam years. If the text is mangled, use your best guess based ONLY on the visible characters in the provided text.
    3. **OUTPUT**: Return ONLY a SINGLE, COMPLETE, VALID JSON ARRAY. No markdown, no preamble.
    4. **STRICT ORDER**: Maintain the EXACT same sequence of questions as they appear in the source.
    5. **COMPLETENESS**: Do not skip any questions. If a question is partially mangled, do your best to reconstruct it.
    6. **LATEX MATH**: 
       - Convert EVERY mathematical symbol, formula, or equation into LaTeX.
       - Use $...$ for inline (e.g., $x^2$) and $$...$$ for standalone formulas.
       - Ensure ALL backslashes are DOUBLE-ESCAPED for JSON (e.g., "\\\\frac", "\\\\sqrt").
       - Example: "The value of 2^3 is" becomes "The value of $2^3$ is".

    ### IMAGE HANDLING ([IMAGE_N] placeholders):
    - The source text has placeholders like [IMAGE_0], [IMAGE_1].
    - Each image belongs to EXACTLY ONE question. An image appears at most once.
    - Assign an `image_index` (integer N) to a question ONLY if:
      a) The placeholder appears immediately before, within, or right after the question text.
      b) The question text explicitly refers to a figure/diagram (e.g., "In the diagram above...").
    - Each `image_index` must be used AT MOST ONCE across all questions. Do NOT assign the same image_index to multiple questions.
    - If no image belongs to a question, set `image_index` to null.

    ### SPECIAL QUESTION TYPES:
    - **Directional/Grouped Questions**: For sections like "Questions 10-15 refer to...", include the shared instruction in the `text` of EVERY question in that group. Use bolding: "**Instruction: Questions 10-15 refer to the diagram.** \\n\\nQuestion: [Actual Question]".
    - **Logical/Statement Questions (GCE Style)**: 
      - Combine the prompt and statements (1, 2, 3...) into the `text` field.
      - Format: "**Instruction: Choose A if 1,2,3 correct...** \\n\\nQuestion: [Prompt] \\n 1. Statement one \\n 2. Statement two..."
      - Replace generic labels (A, B, C, D) with their full definitions if provided in the text.

    ### JSON STRUCTURE:
    Each object MUST have:
    - "text": string (Question text + any relevant instructions + LaTeX math).
    - "marks": integer (default 1 if not specified).
    - "image_index": integer or null.
    - "choices": array of 4 objects:
      - "text": string (Choice text with LaTeX if needed).
      - "is_correct": boolean (exactly one TRUE per question).
    """

    def extract_from_file(self, file_source):
        """
        Main entry point.  `file_source` can be:
          - a local filesystem path (str)
          - a URL string starting with http (downloaded from Supabase)
          - a Django FieldFile / InMemoryUploadedFile with a .name attribute
        """
        if not self.client:
            raise Exception("Gemini client not initialized. Check GEMINI_API_KEY and google-genai installation.")

        # ── Resolve to (bytes, extension) ─────────────────────────────────────
        if hasattr(file_source, 'read'):
            # Django FieldFile or InMemoryUploadedFile
            file_bytes = file_source.read()
            file_name  = getattr(file_source, 'name', '') or ''
            ext        = os.path.splitext(file_name)[1].lower()
        elif isinstance(file_source, str) and file_source.startswith('http'):
            # Supabase public URL — download it
            import urllib.request
            with urllib.request.urlopen(file_source) as resp:
                file_bytes = resp.read()
            ext = os.path.splitext(file_source.split('?')[0])[1].lower()
        elif isinstance(file_source, str):
            # Local path (legacy)
            ext = os.path.splitext(file_source)[1].lower()
            with open(file_source, 'rb') as fh:
                file_bytes = fh.read()
        else:
            raise Exception(f"Unsupported file_source type: {type(file_source)}")

        return self._extract_from_bytes(file_bytes, ext)

    def _extract_from_bytes(self, file_bytes, ext):
        """Process raw bytes and return extracted questions list."""
        import tempfile

        if ext == '.docx':
            if not HAS_DOCX:
                raise Exception("python-docx not installed. Run: pip install python-docx")
            # Write to a temp file so python-docx can open it
            with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
                tmp.write(file_bytes)
                tmp_path = tmp.name
            try:
                text_content, images = self._extract_text_and_images_from_docx(tmp_path)
            finally:
                os.unlink(tmp_path)
        else:
            # PDF — fitz can open from bytes directly
            text_content, images = self._extract_text_from_pdf_bytes(file_bytes)

        if not text_content.strip():
            raise Exception("Could not extract any text from the uploaded file.")

        prompt = f"{self.EXTRACTION_PROMPT}\n\n### SOURCE TEXT:\n{text_content}"

        try:
            response = self.client.models.generate_content(
                model=self.model_name,
                contents=prompt,
            )
            raw_text = response.text.strip()
        except Exception as api_err:
            raise Exception(f"Gemini API error: {api_err}")

        questions_data = self._parse_json(raw_text)

        for q in questions_data:
            idx = q.get('image_index')
            if idx is not None:
                try:
                    idx = int(idx)
                    q['image_data_base64'] = images.get(idx)
                except Exception:
                    q['image_data_base64'] = None
            else:
                q['image_data_base64'] = None

        return questions_data

    def _extract_text_and_images_from_docx(self, file_path):
        doc = Document(file_path)
        text_parts = []
        images = {}
        image_counter = 0

        def get_text_from_element(element):
            parts = []
            for child in element.iter():
                if child.tag.endswith('}t') and child.text:
                    parts.append(child.text)
                if child.tag.endswith('}oMath'):
                    parts.append(" ")
            return "".join(parts).strip()

        for block in doc.element.body:
            if block.tag.endswith('p'):
                for drawing in block.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing'):
                    blip_list = drawing.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
                    for blip in blip_list:
                        rId = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                        if rId and rId in doc.part.related_parts:
                            image_part = doc.part.related_parts[rId]
                            image_base64 = base64.b64encode(image_part.blob).decode('utf-8')
                            mime_type = image_part.content_type or 'image/png'
                            images[image_counter] = f"data:{mime_type};base64,{image_base64}"
                            text_parts.append(f" [IMAGE_{image_counter}] ")
                            image_counter += 1
                p_text = get_text_from_element(block)
                if p_text:
                    text_parts.append(p_text)
            elif block.tag.endswith('tbl'):
                for row in block.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr'):
                    row_cells = []
                    for cell in row.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc'):
                        row_cells.append(get_text_from_element(cell))
                    text_parts.append(" | ".join(row_cells))

        return '\n'.join(text_parts), images

    def _extract_text_from_pdf_bytes(self, file_bytes):
        """Open a PDF from raw bytes via PyMuPDF."""
        doc = fitz.open(stream=file_bytes, filetype='pdf')
        return self._extract_from_fitz(doc)

    def _extract_text_from_pdf(self, file_path):
        doc = fitz.open(file_path)
        return self._extract_from_fitz(doc)

    def _extract_from_fitz(self, doc):
        text_parts = []
        images = {}
        image_counter = 0
        seen_xrefs = {}

        for page_num in range(len(doc)):
            page = doc[page_num]
            elements = []

            blocks = page.get_text("blocks")
            for b in blocks:
                clean_text = b[4].strip()
                if clean_text:
                    elements.append({"type": "text", "y": b[1], "content": b[4]})

            try:
                image_info = page.get_image_info(xrefs=True)
                for info in image_info:
                    xref = info.get('xref', 0)
                    if xref == 0:
                        continue
                    if info.get('width', 0) < 40 or info.get('height', 0) < 40:
                        continue
                    if xref not in seen_xrefs:
                        try:
                            base_image = doc.extract_image(xref)
                            if base_image:
                                image_base64 = base64.b64encode(base_image["image"]).decode('utf-8')
                                image_data = f"data:image/{base_image['ext']};base64,{image_base64}"
                                images[image_counter] = image_data
                                seen_xrefs[xref] = image_counter
                                elements.append({"type": "image", "y": info['bbox'][1], "id": image_counter})
                                image_counter += 1
                        except Exception:
                            continue
            except Exception:
                pass

            elements.sort(key=lambda x: x['y'])
            page_content = []
            for e in elements:
                if e['type'] == 'text':
                    page_content.append(e['content'])
                else:
                    page_content.append(f"\n [IMAGE_{e['id']}] \n")
            text_parts.append("\n".join(page_content))

        doc.close()
        return "\n".join(text_parts), images

    def _repair_json(self, json_str):
        """Repair unescaped LaTeX backslashes for valid JSON."""
        def fix_backslash(match):
            full_match = match.group(0)
            char = match.group(1)
            if char in '"\\/bfnrtu':
                return full_match
            return '\\\\' + char
        json_str = re.sub(r'\\(.)', fix_backslash, json_str, flags=re.DOTALL)
        json_str = re.sub(r',\s*([\]}])', r'\1', json_str)
        return json_str

    def _parse_json(self, raw_text):
        """Cleans and parses the Gemini response into a JSON array."""
        cleaned = re.sub(r'```(?:json)?\s*', '', raw_text).strip().rstrip('`').strip()

        if not cleaned.startswith('['):
            start = cleaned.find('[')
            end = cleaned.rfind(']')
            if start != -1 and end != -1:
                cleaned = cleaned[start:end+1]

        try:
            return json.loads(cleaned)
        except json.JSONDecodeError:
            try:
                repaired = self._repair_json(cleaned)
                return json.loads(repaired)
            except json.JSONDecodeError as e:
                raise Exception(f"JSON Parsing Error: {e}\nContent (first 500 chars): {cleaned[:500]}")

    def save_extracted_quiz(self, quiz_id, extracted_data):
        """Saves questions and options to the database, including extracted images."""
        quiz = Quiz.objects.get(id=quiz_id)
        quiz.questions.all().delete()

        saved_count = 0
        for i, q_data in enumerate(extracted_data):
            # Sanitize: ensure text is always a non-null string
            q_text = q_data.get('text') or ''
            q_text = str(q_text).strip()
            if not q_text:
                print(f"[EXTRACTION] Skipping question {i+1}: empty text")
                continue

            try:
                points = int(q_data.get('marks') or 1)
            except (ValueError, TypeError):
                points = 1

            question = Question.objects.create(
                quiz=quiz,
                text=q_text,
                points=points,
                order=i + 1
            )

            img_b64 = q_data.get('image_data_base64')
            if img_b64 and isinstance(img_b64, str) and img_b64.startswith('data:image'):
                try:
                    format_part, imgstr = img_b64.split(';base64,')
                    ext = format_part.split('/')[-1]
                    data = ContentFile(base64.b64decode(imgstr), name=f'q_{question.id}.{ext}')
                    question.image.save(f'q_{question.id}.{ext}', data, save=True)
                except Exception as img_err:
                    print(f"[EXTRACTION] Error saving image for question {question.id}: {img_err}")

            choices = q_data.get('choices', []) or []
            for j, choice_data in enumerate(choices):
                # Sanitize choice text — skip if null/empty
                choice_text = choice_data.get('text') if choice_data else None
                if choice_text is None:
                    print(f"[EXTRACTION] Skipping null choice {j+1} for question {i+1}")
                    continue
                choice_text = str(choice_text).strip()
                if not choice_text:
                    print(f"[EXTRACTION] Skipping empty choice {j+1} for question {i+1}")
                    continue

                Choice.objects.create(
                    question=question,
                    text=choice_text,
                    is_correct=bool(choice_data.get('is_correct', False))
                )

            saved_count += 1

        quiz.is_extracted = True
        quiz.save()
        print(f"[EXTRACTION] Saved {saved_count}/{len(extracted_data)} questions for quiz {quiz_id}")
